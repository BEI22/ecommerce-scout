"""Excel 导出"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..config import EXPORT_DIR
from ..models import AnalysisReport, Platform, ScoredProduct


def export_report(report: AnalysisReport, filepath: str | None = None) -> Path:
    """导出分析报告到 Excel"""
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    if filepath is None:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filepath = str(EXPORT_DIR / f"选品分析_{report.keyword}_{ts}.xlsx")
    else:
        filepath = str(Path(filepath).resolve())

    wb = Workbook()

    # ── Sheet 1: 选品排名 ──
    ws = wb.active
    ws.title = "选品排名"

    _write_ranking_sheet(ws, report)

    # ── Sheet 2: 市场概览 ──
    ws2 = wb.create_sheet("市场概览")
    _write_overview_sheet(ws2, report)

    # ── Sheet 3: 打分明细 ──
    ws3 = wb.create_sheet("打分明细")
    _write_scoring_sheet(ws3, report)

    # ── Sheet 4: 全部原始数据 ──
    ws4 = wb.create_sheet("原始数据")
    _write_raw_data_sheet(ws4, report)

    wb.save(filepath)
    return Path(filepath)


def _write_ranking_sheet(ws, report: AnalysisReport) -> None:
    """写入选品排名 sheet"""
    # 样式
    header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2B579A", end_color="2B579A", fill_type="solid")
    high_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")     # 绿色：高分
    mid_fill = PatternFill(start_color="FFFDE7", end_color="FFFDE7", fill_type="solid")       # 黄色：中等
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [
        "排名", "综合得分", "商品标题", "平台", "价格(¥)", "月销量",
        "店铺名", "店铺类型", "毛利率%", "跨平台价差(¥)",
        "需求热度", "竞争程度", "利润空间", "平台套利", "趋势判断",
        "选品洞察",
    ]
    col_widths = [6, 10, 45, 10, 12, 12, 18, 12, 10, 14, 10, 10, 10, 10, 10, 55]

    for i, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(i)].width = w

    for row_idx, p in enumerate(report.products, 2):
        platform_name = "淘宝" if p.platform == Platform.TAOBAO else "拼多多"
        values = [
            row_idx - 1,
            p.score,
            p.title,
            platform_name,
            p.price,
            p.sales_count,
            p.shop_name,
            p.shop_type,
            p.estimated_margin_pct,
            p.match_price_diff if p.cross_platform_match else "",
            p.demand_score,
            p.competition_score,
            p.margin_score,
            p.arbitrage_score,
            p.trend_score,
            p.insight,
        ]
        # 行背景
        row_fill = high_fill if p.score >= 70 else mid_fill if p.score >= 50 else None

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = center_align if col_idx != 3 else Alignment(vertical="center", wrap_text=True)
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def _write_overview_sheet(ws, report: AnalysisReport) -> None:
    """写入市场概览 sheet"""
    key_font = Font(name="微软雅黑", bold=True, size=11)
    val_font = Font(name="微软雅黑", size=11)
    title_font = Font(name="微软雅黑", bold=True, size=14, color="2B579A")

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 35

    ws.cell(row=1, column=1, value=f"选品分析报告 — {report.keyword}").font = title_font
    ws.cell(row=2, column=1, value=f"生成时间: {report.generated_at.strftime('%Y-%m-%d %H:%M:%S')}").font = Font(name="微软雅黑", size=10, color="888888")

    data = [
        ("搜索关键词", report.keyword),
        ("", ""),
        ("── 数据概况 ──", ""),
        ("淘宝采集商品数", report.taobao_count),
        ("拼多多采集商品数", report.pdd_count),
        ("有效分析商品数", len(report.products)),
        ("", ""),
        ("── 价格分析 ──", ""),
        ("淘宝均价", f"¥{report.avg_price_taobao:.2f}" if report.avg_price_taobao > 0 else "N/A"),
        ("拼多多均价", f"¥{report.avg_price_pdd:.2f}" if report.avg_price_pdd > 0 else "N/A"),
        ("两平台价差", f"{report.price_gap_pct:+.1f}%"),
        ("", ""),
        ("── 竞争分析 ──", ""),
        ("竞争等级", report.competition_level),
        ("", ""),
        ("── 市场总结 ──", ""),
        ("分析摘要", report.market_summary),
    ]

    for i, (k, v) in enumerate(data, 4):
        ws.cell(row=i, column=1, value=k).font = key_font
        ws.cell(row=i, column=2, value=v).font = val_font


def _write_scoring_sheet(ws, report: AnalysisReport) -> None:
    """写入五维打分明细 sheet"""
    headers = [
        "商品标题", "综合得分", "需求热度(25%)", "竞争程度(25%)",
        "利润空间(30%)", "平台套利(10%)", "趋势判断(10%)",
    ]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)

    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.column_dimensions["A"].width = 45
    for c in range(2, 8):
        ws.column_dimensions[get_column_letter(c)].width = 14

    for row_idx, p in enumerate(report.products, 2):
        values = [
            p.title,
            p.score,
            p.demand_score,
            p.competition_score,
            p.margin_score,
            p.arbitrage_score,
            p.trend_score,
        ]
        for col_idx, val in enumerate(values, 1):
            ws.cell(row=row_idx, column=col_idx, value=val).font = Font(name="微软雅黑", size=10)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def _write_raw_data_sheet(ws, report: AnalysisReport) -> None:
    """写入原始数据"""
    headers = [
        "标题", "平台", "价格", "月销量", "销量原文",
        "店铺名", "店铺类型", "发货地", "标签", "商品链接",
    ]
    header_fill = PatternFill(start_color="808080", end_color="808080", fill_type="solid")
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)

    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["J"].width = 40

    for row_idx, p in enumerate(report.products, 2):
        values = [
            p.title,
            "淘宝" if p.platform == Platform.TAOBAO else "拼多多",
            p.price,
            p.sales_count,
            p.sales_text,
            p.shop_name,
            p.shop_type,
            p.location,
            ", ".join(p.tags),
            p.product_url,
        ]
        for col_idx, val in enumerate(values, 1):
            ws.cell(row=row_idx, column=col_idx, value=val).font = Font(name="微软雅黑", size=10)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
